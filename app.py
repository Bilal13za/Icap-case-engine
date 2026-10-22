"""ICAP Business Case Analysis and Decision Engine — Streamlit UI.

Offline-first study app for ICAP CA students (CAF Business Insights and
Analysis / CFAP Strategic Business Finance, plus their verified 2021
predecessors). Deterministic only: no AI/LLM/cloud calls anywhere.

Run with:  streamlit run app.py
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import sqlite3
import time
from pathlib import Path

import streamlit as st

# ----------------------------------------------------------------------------
# Optional builder modules (graceful degradation if a builder hasn't landed)
# ----------------------------------------------------------------------------
try:
    import dbcore  # noqa: F401

    HAS_DBCORE = True
except Exception:  # pragma: no cover - defensive
    dbcore = None  # type: ignore
    HAS_DBCORE = False

try:
    from engine.rules import Fact, RuleMatch, evaluate_rules, load_rules  # noqa: F401

    HAS_RULES = True
except Exception:
    Fact = RuleMatch = evaluate_rules = load_rules = None  # type: ignore
    HAS_RULES = False

try:
    from engine.frameworks import (  # noqa: F401
        DECISION_TYPES,
        FrameworkRecommendation,
        recommend_frameworks,
    )

    HAS_FRAMEWORKS = True
except Exception:
    DECISION_TYPES = None  # type: ignore
    FrameworkRecommendation = None  # type: ignore
    recommend_frameworks = None  # type: ignore
    HAS_FRAMEWORKS = False

try:
    from engine.recommendations import build_recommendation  # noqa: F401

    HAS_RECOMMENDATIONS = True
except Exception:
    build_recommendation = None  # type: ignore
    HAS_RECOMMENDATIONS = False

try:
    from calcs.calculations import (  # noqa: F401
        REGISTRY,
        CalcResult,
        MissingInputError,
        list_calculations,
        missing_inputs,
        run_calculation,
    )

    HAS_CALCS = True
except Exception:
    REGISTRY = None  # type: ignore
    CalcResult = MissingInputError = None  # type: ignore
    list_calculations = missing_inputs = run_calculation = None  # type: ignore
    HAS_CALCS = False

try:
    from engine import studyaids as _studyaids  # noqa: F401

    HAS_STUDYAIDS = True
except Exception:
    _studyaids = None  # type: ignore
    HAS_STUDYAIDS = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    HAS_OPENPYXL = True
except Exception:  # pragma: no cover - defensive
    Workbook = None  # type: ignore
    HAS_OPENPYXL = False

# ----------------------------------------------------------------------------
# Page config
# ----------------------------------------------------------------------------
st.set_page_config(
    page_title="ICAP Case Analysis & Decision Engine",
    page_icon="📊",
    layout="wide",
)

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"

# ----------------------------------------------------------------------------
# Verification badges
# ----------------------------------------------------------------------------
_BADGE_STYLES = {
    "verified_official": ("#1a7f37", "Verified official"),
    "provisional_predecessor": ("#9a6700", "Provisional — mapped from 2021 predecessor"),
    "supplementary": ("#6e7681", "Supplementary"),
    "unverified": ("#cf222e", "Unverified"),
}


def badge(status: str | None) -> str:
    """Return an HTML badge for a verification_status value."""
    color, label = _BADGE_STYLES.get(
        (status or "").strip(), ("#6e7681", f"Status: {status or 'unknown'}")
    )
    return (
        f'<span style="background:{color}1a;color:{color};border:1px solid {color}55;'
        f'padding:2px 10px;border-radius:12px;font-size:0.78em;font-weight:600;'
        f'white-space:nowrap;">{label}</span>'
    )


def show_badge(status: str | None) -> None:
    st.markdown(badge(status), unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------------
PAPER_CODE_TO_ID = {
    "CAF-7": "PAPER-CAF7-2025",
    "CAF-6": "PAPER-CAF6-2021",
    "CFAP-4-SBF": "PAPER-CFAP4-SBF-2025",
    "CFAP-4-BFD": "PAPER-CFAP4-2021",
}
PAPER_LABEL = {
    "CAF-7": "CAF-7 Business Insights and Analysis (Scheme 2025)",
    "CAF-6": "CAF-6 Managerial and Financial Analysis (Scheme 2021)",
    "CFAP-4-SBF": "CFAP-4 Strategic Business Finance (Scheme 2025)",
    "CFAP-4-BFD": "CFAP-4 Business Finance Decisions (Scheme 2021)",
}

# Controlled fact vocabulary — exactly the fact_keys the rule engine's
# triggers use (engine/rules.py; data/rules.json). A trigger whose fact_key is
# missing is NEVER satisfied: the engine never assumes a fact.
FACT_KEYS = [
    "appraisal_conflict", "budget_variance", "capacity_utilisation",
    "cash_balance", "competitor_price_cut", "current_ratio", "demand",
    "demand_volatility", "dividend_pressure", "ethical_red_flag",
    "exchange_volatility", "fixed_cost_share", "floating_rate_debt",
    "fx_exposure", "gearing", "gross_margin", "interest_rate", "inventory",
    "inventory_days", "investment_opportunities", "largest_customer_share",
    "limiting_factor", "new_entrant", "npv", "operating_cash_flow",
    "operating_profit", "order_delays", "payable_days", "profit_after_tax",
    "project_cost_variance", "receivables_growth_vs_sales", "revenue",
    "rights_issue", "sensitivity_high", "short_term_borrowing",
    "special_order_offer", "strategic_benefit", "supplier_concentration",
    "supplier_quote_below_cost", "synergy_claim", "systems_adequate",
    "tech_disruption",
]
# Trend vocabulary of the rule engine: "up" / "down" / "flat".
# "unknown" is offered for honesty but never matches a trend trigger.
TRENDS = ["up", "down", "flat", "unknown"]

TAG_CLASSES = [
    "quantitative fact", "qualitative fact", "constraint", "assumption",
    "stakeholder concern", "risk indicator", "opportunity",
    "missing information", "external factor", "internal factor",
]

# Canonical list lives in engine/studyaids.py (ANSWER_QUALITY_ITEMS); the
# literal fallback below is byte-identical and only used if that import fails.
SELF_ASSESS_ITEMS = (
    _studyaids.ANSWER_QUALITY_ITEMS
    if HAS_STUDYAIDS
    else [
        "Identified the core business issue / requirement precisely",
        "Applied a relevant framework or technique correctly",
        "Showed all calculations with clear workings",
        "Stated key assumptions explicitly",
        "Considered at least one alternative or sensitivity",
        "Addressed risks and limitations",
        "Reached a clear, justified recommendation / conclusion",
        "Presented the answer in a professional, structured format",
        "Completed within the suggested time",
    ]
)

NAV_ITEMS = [
    "Home",
    "Case Explorer",
    "Scenario Analyzer",
    "Framework Finder",
    "Calculation Workbench",
    "Decision Comparator",
    "Recommendation Builder",
    "Exam Practice",
    "Revision Dashboard",
    "Study Aids",
    "Export",
]

# ----------------------------------------------------------------------------
# Database bootstrap
# ----------------------------------------------------------------------------
DB_OK = False
DB_ERROR = ""
DB_FILE: Path | None = None

if HAS_DBCORE:
    try:
        DB_FILE = Path(dbcore.init_db())
        DB_OK = True
    except Exception as exc:  # pragma: no cover - defensive
        DB_ERROR = str(exc)
else:  # pragma: no cover - defensive temporary fallback
    try:
        fallback_dir = HERE / "app_data"
        fallback_dir.mkdir(exist_ok=True)
        DB_FILE = fallback_dir / "case_engine.db"
        conn = sqlite3.connect(str(DB_FILE))
        conn.close()
        DB_OK = True
        DB_ERROR = "dbcore.py not found — using a temporary local database stub."
    except Exception as exc:
        DB_ERROR = f"dbcore.py not found and fallback failed: {exc}"


def get_conn() -> sqlite3.Connection:
    if HAS_DBCORE and dbcore is not None:
        return dbcore.get_conn()
    if DB_FILE is None:  # pragma: no cover
        raise RuntimeError("No database available.")
    conn = sqlite3.connect(str(DB_FILE))
    conn.row_factory = sqlite3.Row
    return conn


def db_rows(sql: str, params: tuple = ()) -> list[dict]:
    """Run a SELECT and return rows as dicts; [] if the DB is unavailable."""
    if not DB_OK:
        return []
    try:
        conn = get_conn()
        try:
            return [dict(r) for r in conn.execute(sql, params).fetchall()]
        finally:
            conn.close()
    except Exception:
        return []


def db_execute(sql: str, params: tuple = ()) -> None:
    if not DB_OK:
        raise RuntimeError("Database is not available.")
    conn = get_conn()
    try:
        conn.execute(sql, params)
        conn.commit()
    finally:
        conn.close()


def _g(obj, key: str, default=None):
    """Get attribute/key from either a dict or an object."""
    if isinstance(obj, dict):
        return obj.get(key, default)
    return getattr(obj, key, default)


def table_count(table: str) -> int:
    rows = db_rows(f"SELECT COUNT(*) AS n FROM {table}")
    return int(rows[0]["n"]) if rows else 0


def fts_ready(fts_table: str) -> bool:
    try:
        rows = db_rows(
            "SELECT COUNT(*) AS n FROM sqlite_master "
            "WHERE type='table' AND name=? AND sql LIKE '%USING fts5%'",
            (fts_table,),
        )
        if not rows or not rows[0]["n"]:
            return False
        return table_count(fts_table) > 0
    except Exception:
        return False


# ----------------------------------------------------------------------------
# Cached content loaders (DB first, seed JSON fallback)
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def load_papers() -> list[dict]:
    rows = db_rows("SELECT * FROM Papers ORDER BY scheme, code")
    if rows:
        return rows
    p = DATA_DIR / "syllabus.json"
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8")).get("papers", [])
        except Exception:
            return []
    return []


@st.cache_data(show_spinner=False)
def load_questions() -> list[dict]:
    qs = db_rows("SELECT * FROM PracticeQuestions ORDER BY question_id")
    outlines = {r["question_id"]: r for r in db_rows("SELECT * FROM AnswerOutlines")}
    rubrics = {r["question_id"]: r for r in db_rows("SELECT * FROM MarkingRubrics")}
    if not qs:
        p = DATA_DIR / "questions.json"
        if p.exists():
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
                qs = d.get("questions", [])
                outlines = {o["question_id"]: o for o in d.get("answer_outlines", [])}
                rubrics = {r["question_id"]: r for r in d.get("marking_rubrics", [])}
            except Exception:
                qs = []
    for q in qs:
        q["_outline"] = outlines.get(q.get("question_id"), {})
        q["_rubric"] = rubrics.get(q.get("question_id"), {})
    return qs


@st.cache_data(show_spinner=False)
def load_cases() -> tuple[list[dict], dict[str, list[dict]]]:
    cases = db_rows("SELECT * FROM CasePatterns ORDER BY case_id")
    facts = db_rows("SELECT * FROM CaseFacts ORDER BY case_id, fact_id")
    if not cases:
        p = DATA_DIR / "cases.json"
        if p.exists():
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
                cases = d.get("cases", d.get("case_patterns", []))
                facts = d.get("case_facts", [])
            except Exception:
                cases, facts = [], []
    by_case: dict[str, list[dict]] = {}
    for f in facts:
        by_case.setdefault(f.get("case_id", ""), []).append(f)
    return cases, by_case


@st.cache_data(show_spinner=False)
def load_formulas() -> list[dict]:
    rows = db_rows("SELECT * FROM Formulas ORDER BY formula_id")
    if rows:
        return rows
    p = DATA_DIR / "formulas.json"
    if p.exists():
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
            return d.get("formulas", [])
        except Exception:
            return []
    return []


@st.cache_data(show_spinner=False)
def load_frameworks() -> tuple[list[dict], dict[str, list[dict]]]:
    fws = db_rows("SELECT * FROM Frameworks ORDER BY framework_id")
    steps = db_rows("SELECT * FROM FrameworkSteps ORDER BY framework_id, step_no")
    if not fws:
        p = DATA_DIR / "frameworks.json"
        if p.exists():
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
                fws = d.get("frameworks", [])
                steps = d.get("framework_steps", [])
            except Exception:
                fws, steps = [], []
    by_fw: dict[str, list[dict]] = {}
    for s in steps:
        by_fw.setdefault(s.get("framework_id", ""), []).append(s)
    return fws, by_fw


def paper_verification(paper_code: str, papers: list[dict]) -> str | None:
    pid = PAPER_CODE_TO_ID.get(paper_code)
    for p in papers:
        if p.get("paper_id") == pid:
            return p.get("verification_status")
    return None

# ----------------------------------------------------------------------------
# Session state
# ----------------------------------------------------------------------------
def init_state() -> None:
    defaults = {
        "nav": "Home",
        "paper_context": "CAF-7",
        "scenario_facts": [],       # list of {fact_key, value, trend, present}
        "tagged_facts": {},         # fact_id -> classification
        "custom_facts": [],         # free facts added for tagging
        "workings_log": [],         # list of {ts, calc_id, name, inputs, result}
        "selected_frameworks": [],
        "comparison": None,         # dict with options/criteria/scores/note
        "recommendation": None,     # rendered recommendation dict/text
        "reveal_outline": {},       # question_id -> bool
    }
    for k, v in defaults.items():
        st.session_state.setdefault(k, v)


# ----------------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------------
def sidebar() -> None:
    with st.sidebar:
        st.title("📊 ICAP Case Engine")
        st.caption("Business Case Analysis & Decision Engine")
        st.radio("Navigate", NAV_ITEMS, key="nav", label_visibility="collapsed")
        st.divider()
        st.subheader("Module status")
        modules = [
            ("Database (dbcore)", HAS_DBCORE and DB_OK),
            ("Rule engine", HAS_RULES),
            ("Framework engine", HAS_FRAMEWORKS),
            ("Recommendation engine", HAS_RECOMMENDATIONS),
            ("Calculation library", HAS_CALCS),
            ("Excel export (openpyxl)", HAS_OPENPYXL),
        ]
        for name, ok in modules:
            st.markdown(f"{'✅' if ok else '⚠️'} {name}")
        if DB_FILE:
            st.caption(f"DB: {DB_FILE}")
        if DB_ERROR and not (HAS_DBCORE and DB_OK):
            st.warning(DB_ERROR)


# ----------------------------------------------------------------------------
# Screen: Home
# ----------------------------------------------------------------------------
def screen_home() -> None:
    st.header("ICAP Business Case Analysis and Decision Engine")
    st.write(
        "An **offline-first study companion** for ICAP case-style papers. Work a "
        "business case end-to-end: explore case patterns, enter scenario facts, "
        "match rule-based problem indicators, pick analytical frameworks, run "
        "deterministic financial calculations, compare options, build a "
        "structured recommendation, practise exam questions, and track revision."
    )
    st.info(
        "Deterministic by design — this app uses fixed rules, decision tables and "
        "formulas only. It does **not** use AI, the cloud, or the internet, and it "
        "never claims to understand free text: every fact is entered or tagged by "
        "you, manually."
    )

    papers = load_papers()
    st.subheader("Paper context")
    options = list(PAPER_LABEL.keys())
    idx = options.index(st.session_state.paper_context) if st.session_state.paper_context in options else 0
    choice = st.selectbox(
        "Which paper are you studying? (used as the default filter across screens)",
        options, index=idx, format_func=lambda c: PAPER_LABEL[c], key="paper_context",
    )
    for p in papers:
        if p.get("paper_id") == PAPER_CODE_TO_ID.get(choice):
            cols = st.columns([1, 5])
            with cols[0]:
                show_badge(p.get("verification_status"))
            with cols[1]:
                st.caption(
                    f"Source: {p.get('official_source') or '—'}"
                    + (f" · Effective: {p.get('effective_date')}" if p.get("effective_date") else "")
                )
                if p.get("notes"):
                    st.caption(p["notes"])

    st.subheader("⚠️ 2025 scheme rename notice")
    st.warning(
        "**CFAP-4 “Business Finance Decisions” → “Strategic Business Finance”** under the "
        "ICAP Education Scheme 2025, and **CAF-6 “Managerial and Financial Analysis” → "
        "CAF-7 “Business Insights and Analysis”** (verified from the official ICAP FAQ on "
        "the education scheme). The app covers both the fully-verified 2021 syllabuses and "
        "the current 2025 papers; 2025 content mapped from 2021 predecessors is labelled "
        "provisional until ICAP publishes the detailed 2025 syllabus."
    )

    st.subheader("Verification-status legend")
    for status in ("verified_official", "provisional_predecessor", "supplementary", "unverified"):
        st.markdown(badge(status) + f" — {_legend_text(status)}", unsafe_allow_html=True)

    st.subheader("How to work a case here")
    st.markdown(
        "1. **Case Explorer** — find a case pattern and study its facts.\n"
        "2. **Scenario Analyzer** — enter facts, run the rule engine, tag facts manually.\n"
        "3. **Framework Finder** — choose analytical frameworks for your decision type.\n"
        "4. **Calculation Workbench** — run the numbers; save workings.\n"
        "5. **Decision Comparator** — weight criteria, score options.\n"
        "6. **Recommendation Builder** — confirm facts & workings, generate the recommendation.\n"
        "7. **Exam Practice** — timed questions with model outlines and practice rubrics.\n"
        "8. **Revision Dashboard** — attempts, scores and mistakes over time."
    )


def _legend_text(status: str) -> str:
    return {
        "verified_official": "Checked against an official ICAP publication (Directive 1.01 Revised 2024).",
        "provisional_predecessor": "Detailed 2025 syllabus pending official publication; mapped from the verified 2021 predecessor.",
        "supplementary": "Extra study material authored for this app — examinability not verified against ICAP.",
        "unverified": "Claim could not be verified — treat with caution.",
    }[status]


# ----------------------------------------------------------------------------
# Screen: Case Explorer
# ----------------------------------------------------------------------------
_CASE_FIELDS = [
    ("business_context", "Business context"),
    ("industry", "Industry"),
    ("company_size", "Company size"),
    ("main_issue", "Main issue"),
    ("secondary_issues", "Secondary issues"),
    ("decision_to_make", "Decision to make"),
    ("relevant_concepts", "Relevant concepts"),
    ("possible_frameworks", "Possible frameworks"),
    ("quantitative_info", "Quantitative information"),
    ("qualitative_info", "Qualitative information"),
    ("missing_info", "Missing information"),
    ("constraints", "Constraints"),
    ("stakeholders", "Stakeholders"),
    ("alternatives", "Alternatives"),
    ("evaluation_criteria", "Evaluation criteria"),
    ("source_basis", "Source basis"),
    ("topic_ref", "Topic ref"),
]


def _search_cases(query: str) -> list[dict]:
    """FTS5 search with LIKE fallback."""
    if not query.strip():
        return db_rows("SELECT * FROM CasePatterns ORDER BY case_id")
    if fts_ready("case_fts"):
        try:
            terms = " ".join(f'"{w}"' for w in query.split())
            return db_rows(
                "SELECT cp.* FROM case_fts f JOIN CasePatterns cp "
                "ON cp.case_id = f.case_id WHERE case_fts MATCH ? "
                "ORDER BY rank",
                (terms,),
            )
        except Exception:
            pass
    like = f"%{query}%"
    return db_rows(
        "SELECT * FROM CasePatterns WHERE title LIKE ? OR business_context LIKE ? "
        "OR industry LIKE ? OR main_issue LIKE ? ORDER BY case_id",
        (like, like, like, like),
    )


def screen_case_explorer() -> None:
    st.header("Case Explorer")
    cases, facts_by_case = load_cases()
    if not cases:
        st.warning(
            "No case patterns are seeded yet (data/cases.json is still being authored). "
            "This screen will activate automatically once the seed content lands."
        )
        return

    col1, col2, col3 = st.columns([2, 1, 1])
    with col1:
        query = st.text_input("Search cases", placeholder="e.g. working capital textile",
                              key="case_q")
    with col2:
        paper_codes = sorted({c.get("paper_code") for c in cases if c.get("paper_code")})
        paper_f = st.selectbox("Paper", ["All"] + paper_codes,
                               index=(["All"] + paper_codes).index(st.session_state.paper_context)
                               if st.session_state.paper_context in paper_codes else 0)
    with col3:
        industries = sorted({c.get("industry") for c in cases if c.get("industry")})
        ind_f = st.selectbox("Industry", ["All"] + industries)

    results = _search_cases(query) if DB_OK and table_count("CasePatterns") else cases
    if paper_f != "All":
        results = [c for c in results if c.get("paper_code") == paper_f]
    if ind_f != "All":
        results = [c for c in results if c.get("industry") == ind_f]

    st.caption(f"{len(results)} case(s) found"
               + (" · FTS5 full-text search" if fts_ready("case_fts") and query.strip() else ""))
    if not results:
        st.info("No cases match the current filters.")
        return

    labels = {c["case_id"]: f"{c['case_id']} — {c.get('title','')}" for c in results}
    sel_id = st.selectbox("Case", list(labels.keys()), format_func=lambda i: labels[i],
                          key="case_sel")
    case = next(c for c in results if c["case_id"] == sel_id)

    st.subheader(case.get("title", sel_id))
    c1, c2 = st.columns([1, 3])
    with c1:
        show_badge(case.get("verification_status"))
    with c2:
        st.caption(f"{PAPER_LABEL.get(case.get('paper_code'), case.get('paper_code') or '')}"
                   + (f" · {case.get('industry')}" if case.get("industry") else ""))
    if case.get("case_facts"):
        st.markdown(f"**Facts:** {case['case_facts']}")

    with st.expander("All pattern fields", expanded=True):
        for key, label in _CASE_FIELDS:
            val = case.get(key)
            if val:
                st.markdown(f"**{label}:** {val}")

    st.subheader("Case facts")
    facts = facts_by_case.get(sel_id, [])
    if facts:
        for f in facts:
            cols = st.columns([4, 1])
            with cols[0]:
                st.markdown(f"• {f.get('wording') or f.get('normalised')}")
                extra = []
                if f.get("related_issue"):
                    extra.append(f"issue: {f['related_issue']}")
                if f.get("classification"):
                    extra.append(f"class: {f['classification']}")
                if extra:
                    st.caption(" · ".join(extra))
            with cols[1]:
                st.markdown(badge("verified_official")
                            if (f.get("verified_or_assumed") or "").lower().startswith("verif")
                            else badge("supplementary"), unsafe_allow_html=True)
                st.caption(f.get("verified_or_assumed") or "")
    else:
        st.caption("No individual facts recorded for this case yet.")

    st.divider()
    st.subheader("Export this case")
    export_row = {"case_id": case.get("case_id"), "title": case.get("title")}
    for k, _ in _CASE_FIELDS:
        export_row[k] = case.get(k, "")
    st.download_button("⬇️ Download case as CSV", _dicts_to_csv_bytes([export_row]),
                       file_name=f"{sel_id}.csv", mime="text/csv",
                       key=f"dl_case_{sel_id}")

# ----------------------------------------------------------------------------
# Screen: Scenario Analyzer
# ----------------------------------------------------------------------------
def screen_scenario_analyzer() -> None:
    st.header("Scenario Analyzer")
    st.caption(
        "Enter scenario facts in a controlled form, run the deterministic rule "
        "engine, and tag case facts manually. Nothing here interprets free text."
    )

    tab1, tab2 = st.tabs(["Fact entry & rule matching", "Manual fact tagging"])

    with tab1:
        st.subheader("1. Enter facts")
        st.caption("Trend uses the rule engine's vocabulary: **up / down / flat** "
                   "(choose “unknown” rather than guessing). A rule only fires when "
                   "ALL of its triggers match facts you entered — nothing is assumed.")
        with st.form("fact_form", clear_on_submit=True):
            fc1, fc2, fc3, fc4 = st.columns([2, 2, 1, 1])
            with fc1:
                fkey = st.selectbox("Fact key", FACT_KEYS)
            with fc2:
                fval = st.text_input("Value", placeholder="e.g. 68 days, Rs 12m, yes")
            with fc3:
                ftrend = st.selectbox("Trend", TRENDS)
            with fc4:
                fpresent = st.checkbox("Present", value=True)
            if st.form_submit_button("Add fact"):
                st.session_state.scenario_facts.append(
                    {"fact_key": fkey, "value": fval, "trend": ftrend,
                     "present": bool(fpresent)}
                )
                st.success(f"Added: {fkey}")

        facts = st.session_state.scenario_facts
        if facts:
            st.subheader("Entered facts")
            for i, f in enumerate(facts):
                c1, c2, c3, c4, c5 = st.columns([2, 2, 1, 1, 1])
                c1.write(f"**{f['fact_key']}**")
                c2.write(f["value"] or "—")
                c3.write(f["trend"])
                c4.write("yes" if f["present"] else "no")
                if c5.button("✖", key=f"del_fact_{i}"):
                    facts.pop(i)
                    st.rerun()
        else:
            st.info("No facts entered yet. Add facts using the form above.")

        st.subheader("2. Run rule engine")
        if not HAS_RULES:
            st.warning(
                "The rule engine module (engine/rules.py) isn't ready yet — fact "
                "entry and manual tagging above still work. Rule matching will be "
                "enabled automatically once the engine builder's module lands."
            )
        elif not facts:
            st.caption("Add at least one fact to run the rule engine.")
        else:
            if st.button("▶ Run rule matching", type="primary"):
                st.session_state.rule_matches = _run_rules(facts)
            for m in st.session_state.get("rule_matches", []):
                _render_rule_match(m)

    with tab2:
        st.subheader("Classify facts manually")
        st.caption(
            "Manual tagging only — the app never auto-classifies. Your tags are kept "
            "in this session and can be exported as JSON."
        )
        cases, facts_by_case = load_cases()
        options = ["(custom facts)"] + [c["case_id"] for c in cases]
        case_choice = st.selectbox("Case", options, key="tag_case")
        if case_choice == "(custom facts)":
            new_fact = st.text_input("Add a custom fact to tag", key="custom_fact_in")
            if st.button("Add custom fact") and new_fact.strip():
                fid = f"CUSTOM-{len(st.session_state.custom_facts) + 1:03d}"
                st.session_state.custom_facts.append({"fact_id": fid, "wording": new_fact.strip()})
                st.rerun()
            tag_list = st.session_state.custom_facts
        else:
            tag_list = facts_by_case.get(case_choice, [])
        if not tag_list:
            st.info("No facts to tag for this selection.")
        else:
            for f in tag_list:
                fid = f.get("fact_id") or f.get("wording", "")
                cur = st.session_state.tagged_facts.get(fid, "— not tagged —")
                choice = st.selectbox(
                    f.get("wording") or f.get("normalised") or fid,
                    ["— not tagged —"] + TAG_CLASSES,
                    index=(["— not tagged —"] + TAG_CLASSES).index(cur)
                    if cur in (["— not tagged —"] + TAG_CLASSES) else 0,
                    key=f"tag_{fid}",
                )
                if choice != "— not tagged —":
                    st.session_state.tagged_facts[fid] = choice
                elif fid in st.session_state.tagged_facts:
                    del st.session_state.tagged_facts[fid]
            tagged = {k: v for k, v in st.session_state.tagged_facts.items()}
            st.success(f"{len(tagged)} fact(s) tagged in this session.")
            st.download_button(
                "⬇️ Export tagged facts (JSON)",
                json.dumps(tagged, indent=2, ensure_ascii=False).encode("utf-8"),
                file_name="tagged_facts.json", mime="application/json",
                key="dl_tags",
            )


def _run_rules(facts: list[dict]):
    """Build Fact objects and evaluate rules; never raises to the UI."""
    try:
        fact_objs = [Fact(f["fact_key"], f.get("value", ""),
                          f.get("trend", "unknown"), bool(f.get("present", True)))
                     for f in facts]
        rules = load_rules(DATA_DIR / "rules.json")
        return evaluate_rules(fact_objs, rules)
    except Exception as exc:
        st.error(f"Rule engine failed: {exc}")
        return []


def _render_rule_match(m) -> None:
    title = _g(m, "title", _g(m, "rule_id", "Matched rule"))
    rule_id = _g(m, "rule_id", "")
    with st.expander(f"🔎 {title}", expanded=True):
        if rule_id:
            st.caption(f"Rule: {rule_id}")
        interp = _g(m, "interpretation", "")
        if interp:
            st.markdown(f"**Interpretation:** {interp}")
        alt = _g(m, "alternative_explanations", _g(m, "alternatives", ""))
        if alt:
            st.markdown("**Alternative explanations:**")
            _render_bullets(alt)
        info = _g(m, "info_needed", _g(m, "information_needed", ""))
        if info:
            st.markdown("**Information still needed:**")
            _render_bullets(info)
        rec = _g(m, "recommended_analysis", "")
        if rec:
            st.markdown("**Recommended analysis:**")
            _render_bullets(rec)
        concepts = _g(m, "concepts", _g(m, "relevant_concepts", ""))
        if concepts:
            st.caption("Related concepts: " + (
                ", ".join(concepts) if isinstance(concepts, (list, tuple)) else str(concepts)))
        src = _g(m, "source", "")
        if src:
            st.caption(f"Rule source: {src}")
        caution = _g(m, "caution", "")
        st.warning(
            "⚠️ " + (caution or "Indicator only — not a proven cause. Rule out "
                     "alternative explanations before concluding.")
            + " Rule matches flag patterns worth investigating — gather the "
            "information listed above and consider the alternative explanations "
            "before concluding anything."
        )


def _render_bullets(val) -> None:
    if isinstance(val, (list, tuple)):
        for item in val:
            st.markdown(f"- {item}")
    else:
        for line in str(val).splitlines():
            line = line.strip("•*- \t")
            if line:
                st.markdown(f"- {line}")


# ----------------------------------------------------------------------------
# Screen: Framework Finder
# ----------------------------------------------------------------------------


def screen_framework_finder() -> None:
    st.header("Framework Finder")
    st.caption("Recommendations are deterministic: a framework is “ready” only when "
               "every input it needs is among the inputs you tick — the app never "
               "recommends a framework merely because it exists.")
    if not HAS_FRAMEWORKS:
        st.warning(
            "The framework engine (engine/frameworks.py) isn't ready yet. This "
            "screen will activate automatically once the engine builder's module "
            "lands."
        )
        return

    try:
        dt_keys = list(DECISION_TYPES.keys()) if isinstance(DECISION_TYPES, dict) else list(DECISION_TYPES)
    except Exception:
        st.error("Could not read DECISION_TYPES from engine.frameworks.")
        return
    if not dt_keys:
        st.info("No decision types defined yet.")
        return

    dt_key = st.selectbox(
        "Decision type", dt_keys,
        format_func=lambda k: DECISION_TYPES.get(k, k) if isinstance(DECISION_TYPES, dict) else k,
        key="fw_dt",
    )

    if st.button("Find frameworks", type="primary", key="fw_find"):
        try:
            st.session_state.fw_recs = recommend_frameworks(dt_key, []) or []
            st.session_state.fw_inputs_tick = []
        except Exception as exc:
            st.error(f"Framework recommendation failed: {exc}")
    recs = st.session_state.get("fw_recs", [])

    if not recs:
        st.caption("Choose a decision type and click “Find frameworks”.")
        return

    # Step 2: tick the inputs you actually have (union of inputs_needed).
    needed_union: list[str] = []
    for r in recs:
        for need in _g(r, "inputs_needed", []) or []:
            if need not in needed_union:
                needed_union.append(need)
    if needed_union:
        st.subheader("Tick the inputs you have")
        st.caption("Readiness is recomputed from your ticks (case-insensitive substring match).")
        ticked = [n for n in needed_union if st.checkbox(n, key=f"fw_in_{dt_key}_{n}")]
        if st.button("Re-check readiness", key="fw_recheck"):
            try:
                st.session_state.fw_recs = recommend_frameworks(dt_key, ticked) or []
                recs = st.session_state.fw_recs
            except Exception as exc:
                st.error(f"Framework recommendation failed: {exc}")

    st.subheader(f"{len(recs)} recommended framework(s)")
    selected = []
    for i, r in enumerate(recs):
        name = _g(r, "name", _g(r, "framework_id", f"Framework {i+1}"))
        readiness = _g(r, "readiness", "")
        missing = _g(r, "missing", _g(r, "missing_inputs", [])) or []
        ready = (readiness == "ready") or (not missing and readiness != "missing_info")
        with st.expander(f"{'✅' if ready else '🟡'} {name}", expanded=i == 0):
            c1, c2 = st.columns([3, 1])
            with c1:
                if _g(r, "framework_id"):
                    st.caption(f"Framework ID: {_g(r, 'framework_id')}")
                reason = _g(r, "relevance_reason", _g(r, "reason", ""))
                if reason:
                    st.markdown(f"**Why this framework:** {reason}")
                needs = _g(r, "inputs_needed", [])
                if needs:
                    st.markdown("**Inputs needed:**")
                    _render_bullets(needs)
                unsuitable = _g(r, "unsuitable_when", "")
                if unsuitable:
                    st.warning(f"⚠️ Not suitable when: {unsuitable}")
            with c2:
                st.markdown(badge("verified_official") if ready else badge("supplementary"),
                            unsafe_allow_html=True)
                st.caption("Ready to apply" if ready else "Missing info")
                if missing:
                    st.markdown("**Still needed:**")
                    _render_bullets(missing)
            if st.checkbox(f"Select {name}", key=f"fw_sel_{i}_{_g(r, 'framework_id', name)}"):
                selected.append(name)
    if selected:
        st.session_state.selected_frameworks = selected
        st.success(f"Selected: {', '.join(selected)} — carried into the Decision Comparator.")


# ----------------------------------------------------------------------------
# Screen: Calculation Workbench
# ----------------------------------------------------------------------------
def _calc_spec(calc_id: str) -> dict:
    """Full spec from REGISTRY (empty dict if the library isn't loaded)."""
    if REGISTRY is None:
        return {}
    return REGISTRY.get(calc_id, {})


_CATEGORY_LABELS = {
    "breakeven_cvp": "Breakeven / CVP",
    "capital_appraisal": "Capital appraisal",
    "cost_of_capital": "Cost of capital",
    "gearing": "Gearing",
    "liquidity": "Liquidity",
    "performance": "Performance",
    "relevant_costing": "Relevant costing",
    "risk_analysis": "Risk analysis",
    "time_value": "Time value of money",
    "valuation": "Valuation",
    "variances": "Variances",
    "working_capital": "Working capital",
}


def _category_label(cat: str) -> str:
    return _CATEGORY_LABELS.get(cat, cat.replace("_", " ").title())


def _parse_num_list(text: str) -> list[float]:
    vals = []
    for part in text.split(","):
        part = part.strip()
        if part:
            vals.append(float(part))
    return vals


def screen_calc_workbench() -> None:
    st.header("Calculation Workbench")
    st.caption("Deterministic, Decimal-based financial calculations. Every result shows "
               "its formula, substituted steps, interpretation, assumptions and limitations.")

    formulas = load_formulas()
    if formulas:
        with st.expander("📐 Formula library (seed reference)", expanded=False):
            for f in formulas:
                c1, c2 = st.columns([4, 1])
                with c1:
                    st.markdown(f"**{f.get('name') or f.get('formula_id')}**")
                    if f.get("formula_text"):
                        st.code(f["formula_text"])
                    if f.get("assumptions"):
                        st.caption(f"Assumptions: {f['assumptions']}")
                    if f.get("limitations"):
                        st.caption(f"Limitations: {f['limitations']}")
                with c2:
                    show_badge(f.get("verification_status"))

    if not HAS_CALCS:
        st.warning(
            "The calculation library (calcs/calculations.py) isn't ready yet. "
            "This workbench will activate automatically once the calcs builder's "
            "module lands."
        )
        return

    try:
        catalogue = list_calculations() or []
    except Exception as exc:
        st.error(f"Could not list calculations: {exc}")
        return
    if not catalogue:
        st.info("No calculations registered yet.")
        return

    groups: dict[str, list[dict]] = {}
    for entry in catalogue:
        cat = str(entry.get("category", "general"))
        groups.setdefault(_category_label(cat), []).append(entry)
    group = st.selectbox("Group", sorted(groups.keys()), key="calc_group")
    entry = st.selectbox("Calculation", groups[group],
                         format_func=lambda e: e.get("name", e.get("calc_id")),
                         key="calc_sel")
    calc_id = entry.get("calc_id")
    spec = _calc_spec(calc_id)
    tag = entry.get("syllabus_tag", spec.get("syllabus_tag", "supplementary"))

    c1, c2 = st.columns([3, 1])
    with c1:
        st.subheader(entry.get("name", calc_id))
    with c2:
        show_badge("verified_official" if tag == "verified" else "supplementary")
    st.caption(f"Calculation ID: `{calc_id}`")

    required = spec.get("required", []) or []
    optional = spec.get("optional", {}) or {}
    units = spec.get("units", {}) or {}
    num_lists = set(spec.get("numeric_lists", ()) or ())
    text_lists = set(spec.get("text_lists", ()) or ())

    st.subheader("Inputs")
    provided: dict = {}
    list_inputs: dict[str, str] = {}
    with st.form("calc_inputs"):
        for name in required:
            unit = units.get(name, "")
            label = f"{name.replace('_', ' ').title()}" + (f" ({unit})" if unit else "")
            if name in num_lists or name in text_lists:
                list_inputs[name] = st.text_input(
                    label + " — comma-separated",
                    key=f"calc_in_{calc_id}_{name}",
                    help="Enter as a comma-separated list, e.g. 100, 120, 140. Required.",
                    placeholder="100, 120, 140",
                )
            else:
                provided[name] = st.number_input(
                    label, value=0.0, format="%.4f",
                    key=f"calc_in_{calc_id}_{name}", help="Required.",
                )
        for name, default in optional.items():
            unit = units.get(name, "")
            label = f"{name.replace('_', ' ').title()}" + (f" ({unit})" if unit else "")
            try:
                dflt = float(default)
            except (TypeError, ValueError):
                dflt = 0.0
            provided[name] = st.number_input(
                label, value=dflt, format="%.4f",
                key=f"calc_in_{calc_id}_{name}", help="Optional.",
            )
        run = st.form_submit_button("▶ Run calculation", type="primary")

    if run:
        for name, text in list_inputs.items():
            try:
                if name in num_lists:
                    provided[name] = _parse_num_list(text)
                else:
                    provided[name] = [p.strip() for p in text.split(",") if p.strip()]
            except ValueError:
                st.error(f"Input “{name}” must be a comma-separated list of numbers.")
                st.session_state.last_calc = None
                return
        try:
            result = run_calculation(calc_id, **provided)
        except Exception as exc:
            st.session_state.last_calc = None
            _render_calc_error(exc, calc_id, provided)
            return
        # Persist so the result (and its Save button) survive the next rerun.
        st.session_state.last_calc = {
            "calc_id": calc_id,
            "name": entry.get("name", calc_id),
            "spec": spec,
            "provided": dict(provided),
            "result": result,
        }

    lc = st.session_state.get("last_calc")
    if lc and lc.get("calc_id") == calc_id:
        _render_calc_result(lc["name"], lc["spec"], lc["provided"], lc["result"])


def _render_calc_error(exc: Exception, calc_id: str, provided: dict) -> None:
    if MissingInputError is not None and isinstance(exc, MissingInputError):
        missing = list(getattr(exc, "missing", []))
        st.error(f"Cannot compute yet — missing required input(s): {', '.join(missing)}")
        if missing:
            st.markdown("**Missing inputs:**")
            _render_bullets(missing)
    elif "missing" in str(exc).lower():
        st.error(f"Cannot compute yet — {exc}")
        try:
            _render_bullets(missing_inputs(calc_id, provided) if missing_inputs else [])
        except Exception:
            pass
    else:
        st.error(f"Calculation failed: {exc}")
    st.info(
        "📝 What you can still conclude qualitatively: describe the direction "
        "of the effect (e.g. higher gearing → higher financial risk) and note "
        "exactly which figure is needed to quantify it. Never invent the "
        "missing input."
    )


def _render_calc_result(name: str, spec: dict, provided: dict, result) -> None:
    st.subheader("Result")
    formula = _g(result, "formula", spec.get("formula", ""))
    if formula:
        st.markdown("**Formula:**")
        st.code(str(formula))
    steps = _g(result, "steps", "")
    if steps:
        st.markdown("**Substituted steps:**")
        _render_bullets(steps)
    res_val = _g(result, "result", "")
    res_unit = _g(result, "result_unit", "")
    st.metric("Result", f"{res_val}" + (f" {res_unit}" if res_unit else ""))
    interp = _g(result, "interpretation", "")
    if interp:
        st.markdown(f"**Interpretation:** {interp}")
    for label, key in (("Assumptions", "assumptions"), ("Limitations", "limitations")):
        val = _g(result, key, spec.get(key, ""))
        if val:
            st.markdown(f"**{label}:**")
            _render_bullets(val)
    in_units = _g(result, "input_units", spec.get("units", ""))
    if in_units:
        st.caption("Input units: " + ", ".join(
            f"{k} ({v})" for k, v in (in_units.items() if isinstance(in_units, dict) else [])))
    vs = _g(result, "validation_status", "")
    if vs:
        st.caption(f"Validation status: {vs}")

    calc_id = _g(result, "calc_id", "")
    if st.button("💾 Save to workings log", key=f"save_wl_{calc_id}"):
        st.session_state.workings_log.append({
            "ts": dt.datetime.now().isoformat(timespec="seconds"),
            "calc_id": calc_id,
            "name": name,
            "inputs": {k: (", ".join(map(str, v)) if isinstance(v, list) else v)
                       for k, v in provided.items()},
            "result": f"{res_val}" + (f" {res_unit}" if res_unit else ""),
            "interpretation": str(interp or ""),
        })
        st.success("Saved to the workings log (see Export).")
# ----------------------------------------------------------------------------
# Export helpers (CSV via stdlib, formatted Excel via openpyxl)
# ----------------------------------------------------------------------------
def _dicts_to_csv_bytes(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    if not rows:
        return b""
    keys = list(rows[0].keys())
    writer = csv.DictWriter(buf, fieldnames=keys, extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in keys})
    return buf.getvalue().encode("utf-8")


def _dicts_to_excel_bytes(sheets: dict[str, list[dict]]) -> bytes:
    """Multi-sheet workbook: header fill, column widths, borders, frozen panes."""
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31] or "Sheet1")
        if not rows:
            ws["A1"] = "No data"
            continue
        keys = list(rows[0].keys())
        for ci, k in enumerate(keys, start=1):
            cell = ws.cell(row=1, column=ci, value=str(k))
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        for ri, r in enumerate(rows, start=2):
            for ci, k in enumerate(keys, start=1):
                v = r.get(k, "")
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for ci, k in enumerate(keys, start=1):
            maxlen = max([len(str(k))] + [len(str(r.get(k, ""))) for r in rows])
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = min(60, max(12, maxlen + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Screen: Decision Comparator
# ----------------------------------------------------------------------------
OPTION_FIELDS = [
    ("description", "Description"),
    ("financial_impact", "Financial impact"),
    ("nonfinancial_benefits", "Non-financial benefits"),
    ("costs", "Costs"),
    ("risks", "Risks"),
    ("time_requirements", "Time requirements"),
    ("resource_requirements", "Resource requirements"),
    ("stakeholder_effects", "Stakeholder effects"),
    ("strategic_alignment", "Strategic alignment"),
    ("legal_constraints", "Legal constraints"),
    ("assumptions", "Assumptions"),
    ("evidence", "Evidence"),
]
CRITERION_TYPES = ["financial", "qualitative", "strategic", "risk", "feasibility"]


def screen_decision_comparator() -> None:
    st.header("Decision Comparator")
    st.caption("Define options, weight your criteria (weights must sum to 100 — you "
               "enter them; the app never invents them), score each option 1–5.")

    cases, _ = load_cases()
    case_ids = ["(free-form — no case)"] + [c["case_id"] for c in cases]
    case_choice = st.selectbox("Case (optional — prefills options)", case_ids, key="cmp_case")

    if case_choice != "(free-form — no case)" and st.button("Load options from case"):
        opts = db_rows("SELECT * FROM DecisionOptions WHERE case_id = ? ORDER BY option_id",
                       (case_choice,))
        n = st.session_state.get("cmp_n_options", 2)
        for i, o in enumerate(opts[:5]):
            for key, _label in OPTION_FIELDS:
                st.session_state[f"cmp_opt_{i}_{key}"] = o.get(key) or ""
        st.success(f"Loaded {min(len(opts), 5)} option(s) from {case_choice}.")

    n = st.number_input("Number of options", min_value=2, max_value=5, value=2,
                        key="cmp_n_options")
    options: list[dict] = []
    for i in range(int(n)):
        with st.expander(f"Option {i + 1}", expanded=i == 0):
            opt = {"option_id": f"OPT-{i + 1}"}
            for key, label in OPTION_FIELDS:
                opt[key] = st.text_area(label, key=f"cmp_opt_{i}_{key}", height=68)
            if not opt["description"].strip():
                opt["description"] = f"Option {i + 1}"
            options.append(opt)

    st.subheader("Evaluation criteria & weights")
    st.info(
        "Enter each criterion's weight yourself — weights must sum to **100**. "
        "The app will not invent weights for you."
    )
    if "cmp_criteria" not in st.session_state:
        st.session_state.cmp_criteria = [
            {"name": "Financial return", "weight": 40.0, "type": "financial"},
            {"name": "Risk", "weight": 30.0, "type": "risk"},
            {"name": "Strategic fit", "weight": 30.0, "type": "strategic"},
        ]
    criteria = st.session_state.cmp_criteria
    for ci, c in enumerate(criteria):
        c1, c2, c3, c4 = st.columns([3, 1, 2, 1])
        with c1:
            c["name"] = st.text_input("Criterion", value=c["name"], key=f"cmp_crit_{ci}_n")
        with c2:
            c["weight"] = st.number_input("Weight", min_value=0.0, max_value=100.0,
                                          value=float(c["weight"]), step=5.0,
                                          key=f"cmp_crit_{ci}_w")
        with c3:
            c["type"] = st.selectbox("Type", CRITERION_TYPES,
                                     index=CRITERION_TYPES.index(c.get("type", "financial"))
                                     if c.get("type") in CRITERION_TYPES else 0,
                                     key=f"cmp_crit_{ci}_t")
        with c4:
            st.write("")
            if st.button("✖", key=f"cmp_crit_{ci}_del"):
                criteria.pop(ci)
                st.rerun()
    if st.button("＋ Add criterion"):
        criteria.append({"name": f"Criterion {len(criteria) + 1}", "weight": 0.0,
                         "type": "qualitative"})
        st.rerun()

    weight_sum = round(sum(float(c.get("weight") or 0) for c in criteria), 2)
    if abs(weight_sum - 100) > 0.01:
        st.error(f"Weights sum to {weight_sum} — they must sum to 100 before scoring counts.")
        weights_ok = False
    else:
        st.success("Weights sum to 100 ✓")
        weights_ok = True

    st.subheader("Scores (1 = weakest, 5 = strongest)")
    for ci, c in enumerate(criteria):
        st.markdown(f"**{c['name']}** (weight {c['weight']}, {c['type']})")
        cols = st.columns(int(n))
        for i in range(int(n)):
            with cols[i]:
                st.number_input(f"Opt {i + 1}", min_value=1, max_value=5, value=3,
                                key=f"cmp_score_{i}_{ci}")

    totals: list[float] = []
    table: list[dict] = []
    for ci, c in enumerate(criteria):
        row = {"Criterion": c["name"], "Weight": c["weight"], "Type": c["type"]}
        for i in range(int(n)):
            s = int(st.session_state.get(f"cmp_score_{i}_{ci}", 3))
            row[f"Opt {i + 1} score"] = s
            row[f"Opt {i + 1} weighted"] = round(s / 5 * float(c["weight"]), 2)
        table.append(row)
    for i in range(int(n)):
        totals.append(round(sum(r[f"Opt {i + 1} weighted"] for r in table), 2))
    total_row = {"Criterion": "TOTAL", "Weight": weight_sum, "Type": ""}
    for i in range(int(n)):
        total_row[f"Opt {i + 1} score"] = None
        total_row[f"Opt {i + 1} weighted"] = totals[i]
    table.append(total_row)

    st.subheader("Weighted totals")
    st.dataframe(table, width="stretch", hide_index=True)
    if weights_ok and totals:
        best = int(max(range(len(totals)), key=lambda i: totals[i]))
        st.success(f"Highest weighted score: **Option {best + 1}** ({totals[best]} / 100)")

    st.info(
        "💡 **Financial vs qualitative:** financial criteria capture measurable monetary "
        "impact (NPV, cost, cash flow); qualitative criteria capture non-financial factors "
        "(risk, strategic fit, stakeholder effects). Score them separately — do not convert "
        "qualitative scores into rupees. A high total does not override a disqualifying "
        "constraint (e.g. illegality, covenant breach)."
    )
    note = st.text_area("Scenario / sensitivity note",
                        placeholder="e.g. If sales growth is 5% not 10%, Option 2 still wins; "
                        "ranking flips only if Risk weight exceeds 45.",
                        key="cmp_note")
    frameworks_txt = ", ".join(st.session_state.selected_frameworks) if st.session_state.selected_frameworks else ""

    if st.button("💾 Save comparison", type="primary", disabled=not weights_ok):
        st.session_state.comparison = {
            "case_id": None if case_choice.startswith("(free") else case_choice,
            "options": options,
            "criteria": [dict(c) for c in criteria],
            "scores": {f"{i}_{ci}": int(st.session_state.get(f"cmp_score_{i}_{ci}", 3))
                       for i in range(int(n)) for ci in range(len(criteria))},
            "totals": totals,
            "scenario_note": note,
            "frameworks": frameworks_txt,
            "saved_at": dt.datetime.now().isoformat(timespec="seconds"),
        }
        st.success("Comparison saved — pick it up in the Recommendation Builder.")


# ----------------------------------------------------------------------------
# Screen: Recommendation Builder
# ----------------------------------------------------------------------------
def screen_recommendation_builder() -> None:
    st.header("Recommendation Builder")
    st.caption("A recommendation is only generated from facts and workings you "
               "explicitly confirm — nothing is assumed.")
    if not HAS_RECOMMENDATIONS:
        st.warning(
            "The recommendation engine (engine/recommendations.py) isn't ready yet. "
            "This screen will activate automatically once the engine builder's module "
            "lands. Your comparator results and workings log are preserved meanwhile."
        )
        return

    comp = st.session_state.comparison
    if comp and comp.get("options"):
        opt_labels = [f"{o['option_id']}: {(o.get('description') or '')[:60]}"
                      for o in comp["options"]]
        win_idx = st.selectbox("Winning option (from Decision Comparator)",
                               range(len(opt_labels)),
                               format_func=lambda i: opt_labels[i], key="rec_win")
        selected_option_id = comp["options"][win_idx]["option_id"]
        options = comp["options"]
    else:
        st.info("No saved comparison — enter the winning option manually, or go back "
                "to the Decision Comparator.")
        with st.expander("Manual option entry", expanded=True):
            selected_option_id = st.text_input("Winning option ID", value="OPT-1",
                                               key="rec_win_id")
            m_desc = st.text_area("Description", key="rec_m_desc")
            m_fin = st.text_area("Financial impact", key="rec_m_fin")
            m_nfin = st.text_area("Non-financial benefits", key="rec_m_nfin")
            m_costs = st.text_area("Costs", key="rec_m_costs")
            m_risks = st.text_area("Risks", key="rec_m_risks")
            m_assump = st.text_area("Assumptions", key="rec_m_assump")
            m_evid = st.text_area("Evidence", key="rec_m_evid")
        options = [{
            "option_id": selected_option_id or "OPT-1",
            "description": m_desc,
            "financial_impact": m_fin,
            "nonfinancial_benefits": m_nfin,
            "costs": m_costs,
            "risks": m_risks,
            "assumptions": m_assump,
            "evidence": m_evid,
        }]

    st.subheader("1. Confirm facts  (required)")
    fact_choices: list[tuple[str, str]] = []
    for f in st.session_state.scenario_facts:
        fid = f"SCN-{f['fact_key']}"
        fact_choices.append((fid, f"[scenario] {f['fact_key']}: {f.get('value','')} ({f.get('trend')})"))
    for fid, cls in st.session_state.tagged_facts.items():
        fact_choices.append((fid, f"[tagged: {cls}] {fid}"))
    for f in st.session_state.custom_facts:
        fid = f["fact_id"]
        if fid not in st.session_state.tagged_facts:
            fact_choices.append((fid, f"[custom] {f['wording']}"))
    confirmed_fact_ids = [fid for fid, label in fact_choices
                          if st.checkbox(label, key=f"rec_fact_{fid}")]
    if not fact_choices:
        st.caption("No facts available — add some in the Scenario Analyzer first.")

    st.subheader("2. Confirm calculations (optional but recommended)")
    calc_summaries_all = [
        f"{w['name']}: {w['result']}" + (f" — {w['interpretation']}" if w.get("interpretation") else "")
        for w in st.session_state.workings_log
    ]
    calc_summaries = [s for s in calc_summaries_all if st.checkbox(s, key=f"rec_calc_{abs(hash(s)) % 10**8}")]
    if not calc_summaries_all:
        st.caption("Workings log is empty — run the Calculation Workbench first.")

    st.subheader("3. Recommendation inputs")
    st.caption("One item per line for the three list fields.")
    key_assumption = st.text_input("Key assumption", key="rec_assump",
                                   placeholder="e.g. KIBOR stays within 18–22% over the loan term")
    conditions_for_change = st.text_area("Conditions that would change the recommendation "
                                         "(one per line)", key="rec_cond", height=80)
    implementation_steps = st.text_area("Implementation steps (one per line)",
                                        key="rec_impl", height=100)
    kpis = st.text_area("KPIs to monitor (one per line)", key="rec_kpi", height=80)
    review_timing = st.text_input("Review timing", value="Quarterly", key="rec_rev")

    def _lines(text: str) -> list[str]:
        return [l.strip() for l in (text or "").splitlines() if l.strip()]

    can_generate = bool(selected_option_id) and bool(confirmed_fact_ids)
    if not confirmed_fact_ids:
        st.warning("⛔ Generation is blocked until you confirm at least one fact above.")
    if st.button("Generate recommendation", type="primary", disabled=not can_generate):
        try:
            result = build_recommendation(
                selected_option_id=selected_option_id,
                options=options,
                confirmed_fact_ids=confirmed_fact_ids,
                calc_summaries=calc_summaries,
                key_assumption=key_assumption,
                conditions_for_change=_lines(conditions_for_change),
                implementation_steps=_lines(implementation_steps),
                kpis=_lines(kpis),
                review_timing=review_timing,
            )
        except Exception as exc:
            st.error(f"Recommendation engine failed: {exc}")
            return
        text = _recommendation_to_text(result)
        st.session_state.recommendation = {"result": result, "text": text,
                                           "at": dt.datetime.now().isoformat(timespec="seconds")}
        st.success("Recommendation generated.")

    rec = st.session_state.recommendation
    if rec:
        st.divider()
        st.subheader("Structured recommendation")
        _render_recommendation(rec["result"])
        st.download_button("⬇️ Download recommendation (TXT)",
                           rec["text"].encode("utf-8"),
                           file_name="recommendation.txt", mime="text/plain",
                           key="dl_rec_txt")


def _recommendation_to_text(result) -> str:
    if isinstance(result, str):
        return result
    if isinstance(result, dict):
        if result.get("rendered_text"):
            return str(result["rendered_text"])
        lines = []
        if result.get("title"):
            lines.append(str(result["title"]))
            lines.append("")
        sections = result.get("sections")
        if isinstance(sections, dict):
            for k, v in sections.items():
                lines.append(f"== {str(k).replace('_', ' ').upper()} ==")
                if isinstance(v, (list, tuple)):
                    lines.extend(f"- {x}" for x in v)
                else:
                    lines.append(str(v))
                lines.append("")
            return "\n".join(lines)
        for k, v in result.items():
            lines.append(f"== {str(k).replace('_', ' ').upper()} ==")
            if isinstance(v, (list, tuple)):
                lines.extend(f"- {x}" for x in v)
            else:
                lines.append(str(v))
            lines.append("")
        return "\n".join(lines)
    parts = []
    for attr in dir(result):
        if attr.startswith("_"):
            continue
        try:
            v = getattr(result, attr)
        except Exception:
            continue
        if callable(v):
            continue
        parts.append(f"== {attr.replace('_', ' ').upper()} ==\n{v}\n")
    return "\n".join(parts) or str(result)


def _render_recommendation(result) -> None:
    if isinstance(result, str):
        st.markdown(result)
        return
    if isinstance(result, dict) and result.get("title"):
        st.markdown(f"## {result['title']}")
    if isinstance(result, dict) and result.get("rendered_text"):
        st.markdown(str(result["rendered_text"]))
    sections = result.get("sections") if isinstance(result, dict) else None
    items = sections.items() if isinstance(sections, dict) else []
    if not items:
        if isinstance(result, dict):
            items = [(k, v) for k, v in result.items()
                     if k not in ("title", "rendered_text", "sections")]
        else:
            try:
                items = [(a, getattr(result, a)) for a in dir(result)
                         if not a.startswith("_") and not callable(getattr(result, a, None))]
            except Exception:
                items = []
    if not items:
        st.markdown(str(result))
        return
    for k, v in items:
        with st.expander(str(k).replace("_", " ").title(), expanded=False):
            if isinstance(v, (list, tuple)):
                for x in v:
                    st.markdown(f"- {x}")
            else:
                st.markdown(str(v))

# ----------------------------------------------------------------------------
# Screen: Exam Practice
# ----------------------------------------------------------------------------
def screen_exam_practice() -> None:
    st.header("Exam Practice")
    st.caption("Original practice questions written for this app — **not** ICAP past "
               "papers. Model answers are outlines; rubrics are practice rubrics, not "
               "official ICAP marking schemes.")
    questions = load_questions()
    if not questions:
        st.warning("No practice questions available yet.")
        return
    papers = load_papers()

    f1, f2, f3 = st.columns(3)
    with f1:
        pcodes = sorted({q.get("paper_code") for q in questions if q.get("paper_code")})
        pidx = pcodes.index(st.session_state.paper_context) if st.session_state.paper_context in pcodes else 0
        pf = st.selectbox("Paper", pcodes, index=pidx, key="xp_paper",
                          format_func=lambda c: PAPER_LABEL.get(c, c))
    with f2:
        topics = sorted({q.get("topic_ref") for q in questions
                         if q.get("paper_code") == st.session_state.xp_paper and q.get("topic_ref")})
        tf = st.selectbox("Topic", ["All"] + topics, key="xp_topic")
    with f3:
        df = st.selectbox("Difficulty", ["All", "Foundational", "Intermediate", "Advanced"],
                          key="xp_diff")

    filtered = [q for q in questions if q.get("paper_code") == st.session_state.xp_paper]
    if tf != "All":
        filtered = [q for q in filtered if q.get("topic_ref") == tf]
    if df != "All":
        filtered = [q for q in filtered if q.get("difficulty") == df]
    if not filtered:
        st.info("No questions match these filters.")
        return

    qlabels = {q["question_id"]: f"{q['question_id']} · {q['marks']}m · {q['difficulty']} · {q.get('topic_ref','')}"
               for q in filtered}
    qid = st.selectbox("Question", list(qlabels.keys()),
                       format_func=lambda i: qlabels[i], key="xp_q")
    if qid not in qlabels:  # filter changed under a stale selection
        qid = next(iter(qlabels))
    q = next(x for x in filtered if x["question_id"] == qid)

    c1, c2 = st.columns([3, 1])
    with c1:
        st.subheader(f"{qid} — {q.get('marks')} marks · ~{q.get('suggested_minutes')} min")
    with c2:
        show_badge(q.get("verification_status"))
        show_badge(paper_verification(q.get("paper_code"), papers))
    st.markdown(q.get("question_text", ""))
    with st.expander("Question metadata"):
        for label, key in (("Required skills", "required_skills"),
                           ("Expected answer elements", "expected_answer_elements"),
                           ("Calculation requirements", "calculation_requirements"),
                           ("Common mistakes", "common_mistakes")):
            if q.get(key):
                st.markdown(f"**{label}:** {q[key]}")
        st.caption(f"Source: {q.get('source_note') or ''}")

    st.divider()
    col_t1, col_t2 = st.columns([1, 2])
    with col_t1:
        st.subheader("⏱ Timer")
        tkey, skey = f"timer_start_{qid}", f"timer_stop_{qid}"
        b1, b2, b3 = st.columns(3)
        with b1:
            if st.button("Start", key=f"tstart_{qid}"):
                st.session_state[tkey] = time.time()
                st.session_state.pop(skey, None)
        with b2:
            if st.button("Stop", key=f"tstop_{qid}"):
                st.session_state[skey] = time.time()
        with b3:
            if st.button("Refresh", key=f"tref_{qid}"):
                st.rerun()
        start = st.session_state.get(tkey)
        elapsed_min = 0.0
        if start:
            end = st.session_state.get(skey, time.time())
            elapsed = max(0.0, end - start)
            elapsed_min = elapsed / 60
            mm, ss = divmod(int(elapsed), 60)
            st.metric("Elapsed", f"{mm:02d}:{ss:02d}")
            limit = (q.get("suggested_minutes") or 0)
            if limit:
                rem = limit - elapsed / 60
                if rem >= 0:
                    st.caption(f"{rem:.1f} min left of suggested {limit} min")
                else:
                    st.warning(f"{-rem:.1f} min over the suggested {limit} min")
        else:
            st.caption("Timer not started.")
    with col_t2:
        st.subheader("✍️ Answer workspace")
        answer = st.text_area("Your answer", height=220, key=f"ans_{qid}",
                              placeholder="Write your answer here, as you would in the exam…")
        if st.button("📐 Open Calculation Workbench", key=f"wb_{qid}"):
            st.session_state.nav = "Calculation Workbench"
            st.rerun()

    st.divider()
    if st.button("👁 Reveal model answer outline", key=f"rev_{qid}"):
        st.session_state.reveal_outline[qid] = not st.session_state.reveal_outline.get(qid, False)
    if st.session_state.reveal_outline.get(qid):
        outline = q.get("_outline") or {}
        st.markdown("**Model answer outline:**")
        _render_bullets(outline.get("outline_text", "—"))
        st.caption(outline.get("note", "Answer outline — not an official ICAP model answer."))

    rubric = q.get("_rubric") or {}
    rj = rubric.get("rubric_json")
    if isinstance(rj, str):
        try:
            rj = json.loads(rj)
        except Exception:
            rj = []
    with st.expander("📋 Practice marking rubric"):
        if rj:
            for item in rj:
                st.markdown(f"- **{item.get('criterion')}** — {item.get('marks')} marks")
        else:
            st.caption("No rubric detail recorded.")
        st.caption(rubric.get("note", "Practice rubric — not an official ICAP marking scheme."))
        st.caption("This rubric was written for practice only. ICAP's official marking "
                   "schemes are not public; do not treat this as the examiner's scheme.")

    st.divider()
    st.subheader("Self-assessment")
    st.caption("Tick each statement honestly. Score = ticks ÷ 9 × question marks. "
               "This is a practice checklist, not an official assessment.")
    checks = {}
    for i, item in enumerate(SELF_ASSESS_ITEMS):
        checks[item] = st.checkbox(item, key=f"sa_{qid}_{i}")
    score_frac = sum(checks.values()) / len(SELF_ASSESS_ITEMS)
    marks_self = round(score_frac * (q.get("marks") or 0), 1)
    st.metric("Self-assessed score", f"{marks_self} / {q.get('marks')}",
              f"{score_frac:.0%}")

    notes = st.text_input("Attempt notes (optional)", key=f"anotes_{qid}")
    if st.button("💾 Save attempt", type="primary", key=f"save_{qid}"):
        try:
            db_execute(
                "INSERT INTO StudentAttempts (question_id, attempted_at, time_taken_min,"
                " marks_self, checklist_json, answer_text, notes)"
                " VALUES (?,?,?,?,?,?,?)",
                (qid, dt.datetime.now().isoformat(timespec="seconds"),
                 round(elapsed_min, 2), marks_self, json.dumps(checks), answer, notes),
            )
            st.success("Attempt saved — see the Revision Dashboard.")
        except Exception as exc:
            st.error(f"Could not save attempt: {exc}")

    with st.expander("🐞 Log a mistake from this attempt"):
        m1, m2 = st.columns([1, 2])
        with m1:
            mtopic = st.text_input("Topic", value=q.get("topic_ref") or "", key=f"mt_{qid}")
        with m2:
            mdesc = st.text_area("What went wrong?", key=f"md_{qid}", height=80)
        if st.button("Log mistake", key=f"mlog_{qid}"):
            if not mdesc.strip():
                st.warning("Describe the mistake first.")
            else:
                try:
                    last = db_rows(
                        "SELECT attempt_id FROM StudentAttempts WHERE question_id = ?"
                        " ORDER BY attempt_id DESC LIMIT 1", (qid,))
                    aid = last[0]["attempt_id"] if last else None
                    db_execute(
                        "INSERT INTO MistakeLog (attempt_id, topic, description, created_at, resolved)"
                        " VALUES (?,?,?,?,0)",
                        (aid, mtopic.strip(), mdesc.strip(),
                         dt.datetime.now().isoformat(timespec="seconds")),
                    )
                    st.success("Mistake logged.")
                except Exception as exc:
                    st.error(f"Could not log mistake: {exc}")


# ----------------------------------------------------------------------------
# Screen: Revision Dashboard
# ----------------------------------------------------------------------------
def screen_revision_dashboard() -> None:
    st.header("Revision Dashboard")
    attempts = db_rows(
        "SELECT a.*, q.paper_code, q.topic_ref, q.marks, q.difficulty"
        " FROM StudentAttempts a LEFT JOIN PracticeQuestions q"
        " ON q.question_id = a.question_id ORDER BY a.attempted_at DESC"
    )
    mistakes = db_rows("SELECT * FROM MistakeLog ORDER BY created_at DESC")
    if not attempts and not mistakes:
        st.info("No attempts or mistakes logged yet — practise a question in Exam Practice first.")
        return

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        st.metric("Attempts", len(attempts))
    with c2:
        scored = [a for a in attempts if a.get("marks") and a.get("marks_self") is not None]
        avg = (sum(a["marks_self"] / a["marks"] for a in scored) / len(scored)) if scored else 0
        st.metric("Avg self-score", f"{avg:.0%}")
    with c3:
        topics = {a.get("topic_ref") for a in attempts if a.get("topic_ref")}
        topics |= {m.get("topic") for m in mistakes if m.get("topic")}
        st.metric("Topics studied", len(topics))
    with c4:
        unresolved = [m for m in mistakes if not m.get("resolved")]
        st.metric("Open mistakes", len(unresolved))

    st.subheader("Mistakes by topic")
    by_topic = db_rows(
        "SELECT topic, COUNT(*) AS n FROM MistakeLog WHERE resolved = 0"
        " GROUP BY topic ORDER BY n DESC"
    )
    if by_topic:
        st.dataframe(by_topic, width="stretch", hide_index=True)
    else:
        st.caption("No open mistakes. 🎉")

    st.subheader("Repeated mistakes (same topic, ≥2 unresolved)")
    repeated = db_rows(
        "SELECT topic, COUNT(*) AS n FROM MistakeLog WHERE resolved = 0"
        " GROUP BY topic HAVING COUNT(*) >= 2 ORDER BY n DESC"
    )
    if repeated:
        for r in repeated:
            st.warning(f"**{r['topic']}** — {r['n']} unresolved mistakes. Prioritise this topic.")
    else:
        st.caption("None — no topic has 2+ unresolved mistakes.")

    st.subheader("Mistake log")
    if mistakes:
        for m in mistakes:
            with st.expander(f"{'✅' if m.get('resolved') else '🐞'} "
                             f"[{m.get('topic')}] {(m.get('description') or '')[:70]}"):
                st.write(m.get("description"))
                st.caption(f"Logged {m.get('created_at')}")
                if not m.get("resolved") and st.button("Mark resolved",
                                                       key=f"mres_{m['mistake_id']}"):
                    db_execute("UPDATE MistakeLog SET resolved = 1 WHERE mistake_id = ?",
                               (m["mistake_id"],))
                    st.rerun()
    else:
        st.caption("No mistakes logged.")

    st.subheader("Revision history")
    if attempts:
        hist = [{
            "When": a.get("attempted_at"),
            "Question": a.get("question_id"),
            "Paper": a.get("paper_code"),
            "Topic": a.get("topic_ref"),
            "Self-score": (f"{a.get('marks_self')}/{a.get('marks')}"
                           if a.get("marks") else ""),
            "Time (min)": a.get("time_taken_min"),
        } for a in attempts]
        st.dataframe(hist, width="stretch", hide_index=True)
    else:
        st.caption("No attempts yet.")


# ----------------------------------------------------------------------------
# Screen: Study Aids
# ----------------------------------------------------------------------------
def _sa_deck_label(deck: str | None) -> str:
    if HAS_STUDYAIDS:
        return _studyaids.DECK_LABELS.get(deck or "GENERAL", deck or "GENERAL")
    return deck or "GENERAL"


def _sa_mark(card_id: str, status: str) -> None:
    db_execute(
        "INSERT INTO FlashcardProgress(card_id, status, reps, last_seen) "
        "VALUES (?, ?, 1, datetime('now')) "
        "ON CONFLICT(card_id) DO UPDATE SET status=excluded.status, "
        "reps=reps+1, last_seen=excluded.last_seen",
        (card_id, status),
    )
    st.session_state["sa_sess_seen"] = st.session_state.get("sa_sess_seen", 0) + 1
    if status == "known":
        st.session_state["sa_sess_known"] = st.session_state.get("sa_sess_known", 0) + 1


def _sa_advance() -> None:
    st.session_state["sa_idx"] = st.session_state.get("sa_idx", 0) + 1
    st.session_state["sa_show"] = False
    st.rerun()


def _sa_tab_flashcards() -> None:
    cards = db_rows(
        "SELECT c.*, f.name AS framework_name FROM Flashcards c "
        "LEFT JOIN Frameworks f ON f.framework_id = c.framework_id "
        "ORDER BY c.card_id"
    )
    if not cards:
        st.warning("No flashcards found — the database needs re-seeding.")
        return
    type_labels = _studyaids.CARD_TYPE_LABELS

    decks = sorted({c["deck"] or "GENERAL" for c in cards})
    paper = st.selectbox("Deck", decks, format_func=_sa_deck_label, key="sa_deck")
    in_deck = [c for c in cards if (c["deck"] or "GENERAL") == paper]

    fw_names = sorted({c["framework_name"] or "General" for c in in_deck})
    fw_choice = st.selectbox("Framework", ["All"] + fw_names, key="sa_fw")
    types = sorted({c["card_type"] for c in in_deck})
    sel_types = st.multiselect(
        "Card types", options=types, default=types,
        format_func=lambda t: type_labels.get(t, t), key="sa_types")
    deck = [c for c in in_deck
            if (fw_choice == "All"
                or (c["framework_name"] or "General") == fw_choice)
            and c["card_type"] in sel_types]
    if not deck:
        st.info("No cards match this selection.")
        return

    prog = {r["card_id"]: r for r in
            db_rows("SELECT card_id, status, reps FROM FlashcardProgress")}
    known = sum(1 for c in deck
                if prog.get(c["card_id"], {}).get("status") == "known")
    st.caption(f"Deck progress: {known}/{len(deck)} known "
               f"(persisted across sessions)")
    st.progress(known / len(deck))

    if st.button("🔀 Shuffle deck", key="sa_shuffle"):
        import random  # noqa: PLC0415

        order = [c["card_id"] for c in deck]
        random.shuffle(order)
        st.session_state["sa_order"] = order
        st.session_state["sa_idx"] = 0
        st.session_state["sa_show"] = False

    ids = [c["card_id"] for c in deck]
    order = st.session_state.get("sa_order")
    if not order or set(order) != set(ids):
        order = ids
        st.session_state["sa_order"] = order
        st.session_state["sa_idx"] = 0
    idx = st.session_state.get("sa_idx", 0) % len(order)
    by_id = {c["card_id"]: c for c in deck}
    card = by_id[order[idx]]

    st.subheader(f"Card {idx + 1} of {len(order)}")
    show_badge(card.get("verification_status"))
    st.markdown(
        f"**{type_labels.get(card['card_type'], card['card_type'])}**"
        + (f" · {card['framework_name']}" if card.get("framework_name") else "")
    )
    with st.container(border=True):
        st.markdown(card["front"])
    if st.button("Show answer", key="sa_show_btn"):
        st.session_state["sa_show"] = True
    if st.session_state.get("sa_show"):
        with st.container(border=True):
            st.markdown(card["back"])
        c1, c2, c3 = st.columns(3)
        if c1.button("✅ Know", key="sa_know"):
            _sa_mark(card["card_id"], "known")
            _sa_advance()
        if c2.button("🤔 Unsure", key="sa_unsure"):
            _sa_mark(card["card_id"], "unsure")
            _sa_advance()
        if c3.button("⏭ Skip", key="sa_skip"):
            _sa_advance()
    seen_n = st.session_state.get("sa_sess_seen", 0)
    if seen_n:
        known_n = st.session_state.get("sa_sess_known", 0)
        st.caption(f"This session: {known_n}/{seen_n} marked known")


def _sa_tab_mindmaps() -> None:
    maps = db_rows(
        "SELECT m.*, f.name AS framework_name FROM MindMaps m "
        "JOIN Frameworks f ON f.framework_id = m.framework_id "
        "ORDER BY f.name"
    )
    if not maps:
        st.warning("No mind maps found — the database needs re-seeding.")
        return
    choice = st.selectbox("Mind map", [m["title"] for m in maps], key="sa_map")
    m = next(x for x in maps if x["title"] == choice)
    fw = db_rows("SELECT * FROM Frameworks WHERE framework_id = ?",
                (m["framework_id"],))
    steps = db_rows("SELECT * FROM FrameworkSteps WHERE framework_id = ? "
                    "ORDER BY step_no", (m["framework_id"],))
    if not fw:
        st.warning("Framework data missing for this map.")
        return
    try:
        order = json.loads(m["branch_order_json"]) if m["branch_order_json"] else None
    except Exception:
        order = None
    html_doc = _studyaids.mindmap_html(dict(fw[0]),
                                       [dict(s) for s in steps], order)
    st.components.v1.html(html_doc, height=720, scrolling=True)


def _sa_tab_formulas() -> None:
    formulas = {r["formula_id"]: r for r in db_rows("SELECT * FROM Formulas")}
    groups = _studyaids.FORMULA_SHEET_GROUPS
    gsel = st.selectbox("Quick sheet", [g[0] for g in groups], key="sa_sheet")
    title, paper, fids = next(g for g in groups if g[0] == gsel)
    st.caption(f"Primary paper grouping: {_sa_deck_label(paper)} — "
               f"many formulas apply to both papers.")
    for fid in fids:
        f = formulas.get(fid)
        if not f:
            continue
        with st.expander(f["name"]):
            show_badge(f.get("verification_status"))
            st.markdown(f"**Formula:** {f['formula_text']}")
            try:
                inps = json.loads(f["inputs_json"]) if f.get("inputs_json") else []
            except Exception:
                inps = []
            if inps:
                st.markdown("**Inputs:**")
                for i in inps:
                    st.markdown(f"- `{i.get('name')}` ({i.get('unit', '')}) — "
                                f"{i.get('description', '')}")
            if f.get("assumptions"):
                st.markdown(f"**Assumptions:** {f['assumptions']}")
            if f.get("limitations"):
                st.markdown(f"**Limitations:** {f['limitations']}")
            st.caption(f"Validation: {f.get('validation_status', 'unknown')}")


def _sa_tab_checklists() -> None:
    st.subheader("Recommendation checklist")
    st.caption("The 11 sections of the app's recommendation template — tick each "
               "off as you complete it in the Recommendation Builder.")
    rec = _studyaids.RECOMMENDATION_CHECKLIST
    done_r = 0
    for key, label in rec:
        k = f"sa_chk_rec_{key}"
        if k not in st.session_state:
            st.session_state[k] = False
        if st.checkbox(label, key=k):
            done_r += 1
    st.progress(done_r / len(rec), text=f"{done_r}/{len(rec)} complete")

    st.subheader("Exam answer-quality checklist")
    st.caption("The same 9-point self-assessment used in Exam Practice.")
    items = _studyaids.ANSWER_QUALITY_ITEMS
    done_q = 0
    for i, label in enumerate(items):
        k = f"sa_chk_q_{i}"
        if k not in st.session_state:
            st.session_state[k] = False
        if st.checkbox(label, key=k):
            done_q += 1
    st.progress(done_q / len(items), text=f"{done_q}/{len(items)} complete")

    if st.button("Reset checklists", key="sa_chk_reset"):
        for key, _ in rec:
            st.session_state[f"sa_chk_rec_{key}"] = False
        for i in range(len(items)):
            st.session_state[f"sa_chk_q_{i}"] = False
        st.rerun()


def screen_study_aids() -> None:
    st.header("📚 Study Aids")
    st.caption("Flashcards, mind maps, formula sheets and checklists — generated "
               "deterministically from the app's own syllabus, framework, formula "
               "and rule content. Everything runs offline on your machine.")
    if not HAS_STUDYAIDS:
        st.warning("Study-aids engine unavailable.")
        return
    tabs = st.tabs(["Flashcards", "Mind Maps", "Formula Sheets", "Checklists"])
    with tabs[0]:
        _sa_tab_flashcards()
    with tabs[1]:
        _sa_tab_mindmaps()
    with tabs[2]:
        _sa_tab_formulas()
    with tabs[3]:
        _sa_tab_checklists()


# ----------------------------------------------------------------------------
# Screen: Export
# ----------------------------------------------------------------------------
def screen_export() -> None:
    st.header("Export")
    st.caption("CSV via the standard library; Excel workbooks are formatted with "
               "header fills, column widths, borders and frozen panes (openpyxl).")

    def dl_csv(label: str, rows: list[dict], fname: str, key: str) -> None:
        st.download_button(f"⬇️ {label} (CSV)", _dicts_to_csv_bytes(rows),
                           file_name=fname, mime="text/csv", key=key,
                           disabled=not rows)

    # Workings log
    st.subheader("Calculation workings log")
    wl = st.session_state.workings_log
    wl_rows = [{**w, "inputs": json.dumps(w.get("inputs", {}), ensure_ascii=False)} for w in wl]
    dl_csv("Workings log", wl_rows, "workings_log.csv", "exp_wl_csv")
    if wl and HAS_OPENPYXL:
        st.download_button("⬇️ Workings log (Excel)",
                           _dicts_to_excel_bytes({"Workings": wl_rows}),
                           file_name="workings_log.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           key="exp_wl_xls")

    # Decision comparison
    st.subheader("Decision comparison")
    comp = st.session_state.comparison
    if comp:
        opt_rows = [dict(o) for o in comp.get("options", [])]
        crit_rows = []
        for ci, c in enumerate(comp.get("criteria", [])):
            row = {"criterion": c["name"], "weight": c["weight"], "type": c["type"]}
            for i in range(len(comp.get("options", []))):
                row[f"option_{i+1}_score"] = comp.get("scores", {}).get(f"{i}_{ci}", "")
            crit_rows.append(row)
        dl_csv("Comparison — options", opt_rows, "comparison_options.csv", "exp_co_csv")
        dl_csv("Comparison — criteria & scores", crit_rows, "comparison_criteria.csv", "exp_cc_csv")
        if HAS_OPENPYXL:
            st.download_button(
                "⬇️ Comparison (Excel)",
                _dicts_to_excel_bytes({
                    "Options": opt_rows,
                    "CriteriaAndScores": crit_rows,
                    "Meta": [{"key": k, "value": str(v)}
                             for k, v in comp.items()
                             if k not in ("options", "criteria", "scores")],
                }),
                file_name="decision_comparison.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="exp_c_xls")
    else:
        st.caption("No comparison saved yet.")

    # Recommendation
    st.subheader("Recommendation")
    rec = st.session_state.recommendation
    if rec:
        res = rec["result"]
        sections = res.get("sections") if isinstance(res, dict) else None
        if isinstance(sections, dict):
            sec_rows = [{"section": str(k).replace("_", " ").title(),
                         "content": (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else str(v))}
                        for k, v in sections.items()]
        elif isinstance(res, dict):
            sec_rows = [{"section": str(k).replace("_", " ").title(),
                         "content": (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else str(v))}
                        for k, v in res.items() if k != "rendered_text"]
        else:
            sec_rows = [{"section": "recommendation", "content": _recommendation_to_text(res)}]
        dl_csv("Recommendation sections", sec_rows, "recommendation.csv", "exp_r_csv")
        if HAS_OPENPYXL:
            st.download_button(
                "⬇️ Recommendation (Excel)",
                _dicts_to_excel_bytes({"Recommendation": sec_rows}),
                file_name="recommendation.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="exp_r_xls")
    else:
        st.caption("No recommendation generated yet.")

    # Attempts & mistakes
    st.subheader("Attempts & mistakes")
    attempts = db_rows("SELECT * FROM StudentAttempts ORDER BY attempted_at")
    mistakes = db_rows("SELECT * FROM MistakeLog ORDER BY created_at")
    dl_csv("Attempts", attempts, "attempts.csv", "exp_a_csv")
    dl_csv("Mistakes", mistakes, "mistakes.csv", "exp_m_csv")
    if HAS_OPENPYXL and (attempts or mistakes):
        st.download_button(
            "⬇️ Attempts & mistakes (Excel)",
            _dicts_to_excel_bytes({"Attempts": attempts, "Mistakes": mistakes}),
            file_name="attempts_mistakes.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="exp_am_xls")

    # Tagged facts
    st.subheader("Tagged facts")
    tagged = st.session_state.tagged_facts
    st.download_button("⬇️ Tagged facts (JSON)",
                       json.dumps(tagged, indent=2, ensure_ascii=False).encode("utf-8"),
                       file_name="tagged_facts.json", mime="application/json",
                       key="exp_t_json", disabled=not tagged)

    # Combined workbook
    st.subheader("Everything")
    if HAS_OPENPYXL:
        if st.button("⬇️ Build combined Excel workbook", key="exp_all"):
            sheets: dict[str, list[dict]] = {}
            if wl_rows:
                sheets["Workings"] = wl_rows
            if comp:
                sheets["CompOptions"] = [dict(o) for o in comp.get("options", [])]
            if rec:
                rres = rec["result"]
                rsec = rres.get("sections") if isinstance(rres, dict) else None
                if isinstance(rsec, dict):
                    sheets["Recommendation"] = [
                        {"section": str(k).replace("_", " ").title(),
                         "content": (", ".join(map(str, v)) if isinstance(v, (list, tuple)) else str(v))}
                        for k, v in rsec.items()]
                else:
                    sheets["Recommendation"] = [
                        {"section": "recommendation", "content": rec["text"]}]
            if attempts:
                sheets["Attempts"] = attempts
            if mistakes:
                sheets["Mistakes"] = mistakes
            if not sheets:
                st.warning("Nothing to export yet.")
                st.session_state.combined_xls = None
            else:
                st.session_state.combined_xls = _dicts_to_excel_bytes(sheets)
                st.success(f"Workbook built with sheets: {', '.join(sheets.keys())}")
        if st.session_state.get("combined_xls"):
            st.download_button(
                "⬇️ Download case_engine_export.xlsx",
                st.session_state.combined_xls,
                file_name="case_engine_export.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="exp_all_xls")


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main() -> None:
    init_state()
    if not DB_OK:
        st.error(
            "Database initialisation failed: "
            + (DB_ERROR or "unknown error")
            + ". Content screens that need the database are unavailable; "
            "Exam Practice can still use the bundled seed JSON."
        )
    sidebar()
    screens = {
        "Home": screen_home,
        "Case Explorer": screen_case_explorer,
        "Scenario Analyzer": screen_scenario_analyzer,
        "Framework Finder": screen_framework_finder,
        "Calculation Workbench": screen_calc_workbench,
        "Decision Comparator": screen_decision_comparator,
        "Recommendation Builder": screen_recommendation_builder,
        "Exam Practice": screen_exam_practice,
        "Revision Dashboard": screen_revision_dashboard,
        "Study Aids": screen_study_aids,
        "Export": screen_export,
    }
    screens.get(st.session_state.nav, screen_home)()


if __name__ == "__main__":
    main()
